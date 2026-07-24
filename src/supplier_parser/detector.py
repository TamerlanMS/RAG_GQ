"""
SupplierDetector — определяет поставщика по сигнатуре файла.
Сигнатура: названия листов + ключевые слова в первых строках.
НЕ зависит от имени файла.
"""
from __future__ import annotations

from typing import Optional

import openpyxl


def detect_supplier(file_path: str) -> Optional[str]:
    """
    Возвращает supplier_code (himel/dkc/iek/ekf/schneider/legrand)
    или None если поставщик не распознан.
    """
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheets = [s.lower() for s in wb.sheetnames]
        wb.close()
    except Exception:
        return None

    # Legrand: лист «Тариф» + проверяем что есть LG- артикулы
    if "тариф" in sheets and "удаленные" in sheets:
        if _check_cell_contains(file_path, "Тариф", 5, 0, "LG-"):
            return "legrand"

    # Himel: лист «Тариф» без «Удаленные», есть «Референс» в заголовке
    if sheets == ["тариф"] or (len(sheets) == 1 and sheets[0] == "тариф"):
        if _check_cell_contains(file_path, "Тариф", 3, 0, "Референс"):
            return "himel"

    # Schneider: лист «Тариф 2026» (с «Тариф» в имени, строка 1 содержит «Коллекция»)
    for s in wb_sheetnames(file_path):
        if "тариф" in s.lower():
            if _check_cell_contains(file_path, s, 1, 0, "Коллекция"):
                return "schneider"

    # DKC: лист «Прайс ДКС» или «Запрос_1»
    if "прайс дкс" in sheets or "запрос_1" in sheets:
        return "dkc"

    # IEK: лист «Прайс» с заголовком «Артикул» на строке 7
    if "прайс" in sheets:
        if _check_cell_contains(file_path, _find_sheet(file_path, "прайс"), 7, 0, "Артикул"):
            return "iek"

    # EKF: лист «Продукция EKF»
    if "продукция ekf" in sheets:
        return "ekf"

    return None


def wb_sheetnames(file_path: str):
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        names = wb.sheetnames[:]
        wb.close()
        return names
    except Exception:
        return []


def _find_sheet(file_path: str, keyword: str) -> str:
    """Найти лист по ключевому слову (без учёта регистра)."""
    for s in wb_sheetnames(file_path):
        if keyword.lower() in s.lower():
            return s
    return ""


def _check_cell_contains(file_path: str, sheet: str, row: int, col: int, text: str) -> bool:
    """Проверить что ячейка [row, col] содержит text."""
    if not sheet:
        return False
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        if sheet not in wb.sheetnames:
            wb.close()
            return False
        ws = wb[sheet]
        for r in ws.iter_rows(min_row=row, max_row=row, values_only=True):
            cell = r[col] if len(r) > col else None
            wb.close()
            return cell is not None and text.lower() in str(cell).lower()
        wb.close()
    except Exception:
        pass
    return False

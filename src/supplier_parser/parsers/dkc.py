"""
Парсер прайса DKC.
Используем лист «Запрос_1» — он содержит плоские данные без иерархии.

Колонки (0-based):
  0  Материал       = article
  1  Остаток/Назв.  = name
  2  '#' для товаров (категории не имеют '#')
  3  БазЕдиницаИзмерения = unit
  9  Цена с НДС
  10 Цена без НДС  = price_base

Строка 1 — заголовок, данные с строки 2.
Фильтр: col[2] == '#' (реальный товар, не категория).
"""
from __future__ import annotations

from typing import List

import openpyxl

from src.supplier_parser.models import SupplierProduct
from src.supplier_parser.normalizer import clean_str, clean_price, clean_unit
from src.supplier_parser.parsers.base import BaseParser


class DKCParser(BaseParser):
    supplier_code = "dkc"
    SHEET = "Запрос_1"
    DATA_FROM = 2   # 1-based (row 1 = header)

    def parse(self, file_path: str) -> List[SupplierProduct]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb[self.SHEET]

        # Дата из имени колонки заголовка (формат "20260326 Цена c НДС...")
        price_date = None
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            for cell in row:
                if cell and str(cell).strip()[:8].isdigit():
                    import re
                    m = re.match(r"(\d{8})", str(cell))
                    if m:
                        try:
                            from datetime import datetime
                            price_date = datetime.strptime(m.group(1), "%Y%m%d").date()
                        except ValueError:
                            pass
            break

        items: List[SupplierProduct] = []
        for row in ws.iter_rows(min_row=self.DATA_FROM, values_only=True):
            # Только реальные товары: маркер '#' в col2
            if row[2] != "#":
                continue
            article = clean_str(row[0])
            name = clean_str(row[1])
            if not article or not name:
                continue
            price_base = clean_price(row[10])
            price_with_vat = clean_price(row[9])
            if price_base is None and price_with_vat is None:
                continue
            if price_base is None and price_with_vat is not None:
                from src.supplier_parser.normalizer import price_excl_vat
                price_base = price_excl_vat(price_with_vat)
            try:
                items.append(SupplierProduct(
                    supplier_code=self.supplier_code,
                    article=article,
                    name=name,
                    unit=clean_unit(row[3]),
                    price_base=price_base,
                    price_with_vat=price_with_vat,
                    brand="DKC",
                    price_date=price_date,
                ))
            except Exception:
                continue

        wb.close()
        return self._dedupe(items)

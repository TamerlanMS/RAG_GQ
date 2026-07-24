"""
Парсер прайса EKF.
Листы: «Продукция EKF» и «Новинки» (одинаковая структура).
Заголовок: строка 12 (1-based)
Данные: строка 13+

Колонки (0-based):
  0  A  Артикул
  1  B  Номенклатура
  5  F  Ед.изм.
  6  G  Базовая цена с НДС, KZT    → price_with_vat
  9  J  Цена с учётом скидок с НДС → price_base (используем как ближайшую к закупочной)
  10 K  Мин. норма отпуска
  11 L  Кол-во в упаковке → pack_qty
"""
from __future__ import annotations

from typing import List

import openpyxl

from src.supplier_parser.models import SupplierProduct
from src.supplier_parser.normalizer import (
    clean_str, clean_price, clean_unit, clean_int, price_excl_vat,
)
from src.supplier_parser.parsers.base import BaseParser

_SHEETS = ["Продукция EKF", "Новинки"]
_HEADER_ROW = 12
_DATA_FROM = 13


class EKFParser(BaseParser):
    supplier_code = "ekf"

    def parse(self, file_path: str) -> List[SupplierProduct]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

        # Дата из B2
        price_date = None
        for sheet_name in _SHEETS:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows(min_row=2, max_row=2, values_only=True):
                    cell = row[1] if len(row) > 1 else None
                    if cell and hasattr(cell, "year"):
                        price_date = cell.date() if hasattr(cell, "date") else None
                    break
                break

        all_items: List[SupplierProduct] = []
        for sheet_name in _SHEETS:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=_DATA_FROM, values_only=True):
                article = clean_str(row[0])
                name = clean_str(row[1])
                if not article or not name:
                    continue
                price_with_vat = clean_price(row[6]) if len(row) > 6 else None
                price_discount = clean_price(row[9]) if len(row) > 9 else None
                if price_with_vat is None:
                    continue
                # price_base = цена без НДС (используем скидочную если есть)
                base_vat = price_discount if price_discount else price_with_vat
                price_base = price_excl_vat(base_vat)
                try:
                    all_items.append(SupplierProduct(
                        supplier_code=self.supplier_code,
                        article=article,
                        name=name,
                        unit=clean_unit(row[5]) if len(row) > 5 else "шт",
                        price_base=price_base,
                        price_with_vat=price_with_vat,
                        pack_qty=clean_int(row[11]) if len(row) > 11 else None,
                        brand="EKF",
                        price_date=price_date,
                    ))
                except Exception:
                    continue

        wb.close()
        return self._dedupe(all_items)

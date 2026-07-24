"""
Парсер прайса Himel.
Лист: «Тариф»
Заголовок: строка 3 (0-based row 2)
Данные: строка 4+

Колонки:
  A=0  Референс (артикул)
  B=1  Описание (наименование)
  C=2  Единица измерения
  D=3  Кратность (pack_qty)
  E=4  Тариф без НДС
  F=5  Комментарий
  G=6  NTIN
"""
from __future__ import annotations

from datetime import date
from typing import List

import openpyxl

from src.supplier_parser.models import SupplierProduct
from src.supplier_parser.normalizer import clean_str, clean_price, clean_unit, clean_int, price_incl_vat
from src.supplier_parser.parsers.base import BaseParser


class HimelParser(BaseParser):
    supplier_code = "himel"
    SHEET = "Тариф"
    HEADER_ROW = 3   # 1-based
    DATA_FROM = 4    # 1-based

    def parse(self, file_path: str) -> List[SupplierProduct]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb[self.SHEET]

        # Определяем дату прайса из заголовка (ищем дату в строке 3)
        price_date = None
        for row in ws.iter_rows(min_row=self.HEADER_ROW, max_row=self.HEADER_ROW, values_only=True):
            for cell in row:
                if cell and "тариф" in str(cell).lower():
                    import re
                    m = re.search(r"(\d{2}\.\d{2}\.\d{4})", str(cell))
                    if m:
                        try:
                            from datetime import datetime
                            price_date = datetime.strptime(m.group(1), "%d.%m.%Y").date()
                        except ValueError:
                            pass

        items: List[SupplierProduct] = []
        for row in ws.iter_rows(min_row=self.DATA_FROM, values_only=True):
            article = clean_str(row[0])
            name = clean_str(row[1])
            if not article or not name:
                continue
            price_base = clean_price(row[4])
            if price_base is None:
                continue
            try:
                items.append(SupplierProduct(
                    supplier_code=self.supplier_code,
                    article=article,
                    name=name,
                    unit=clean_unit(row[2]),
                    price_base=price_base,
                    price_with_vat=price_incl_vat(price_base),
                    pack_qty=clean_int(row[3]),
                    ntin=clean_str(row[6]),
                    price_date=price_date,
                ))
            except Exception:
                continue

        wb.close()
        return self._dedupe(items)

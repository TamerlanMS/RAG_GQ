"""
Парсер прайса IEK / ITK.
Лист: «Прайс»
Заголовок: строка 7 (1-based)
Данные: строка 9+ (строка 8 пустая)

Колонки (0-based):
  0  A  Артикул
  1  B  Наименование
  3  D  Категория (уровень 1)
  4  E  Группа (уровень 2)
  5  F  Подгруппа (уровень 3)
  7  H  Ед.изм.
  8  I  NTIN
  9  J  Статус
  11 L  Кратность (pack_qty)
  14 O  Базовая цена с НДС → price_with_vat
  16 Q  Рекомендованная оптовая цена с НДС → price_base
"""
from __future__ import annotations

from datetime import date
from typing import List

import openpyxl

from src.supplier_parser.models import SupplierProduct
from src.supplier_parser.normalizer import (
    clean_str, clean_price, clean_unit, clean_int, price_excl_vat,
)
from src.supplier_parser.parsers.base import BaseParser


class IEKParser(BaseParser):
    supplier_code = "iek"
    SHEET = "Прайс"
    HEADER_ROW = 7
    DATA_FROM = 9

    def parse(self, file_path: str) -> List[SupplierProduct]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb[self.SHEET]

        # Дата из строки 2 (ячейка B2 — datetime)
        price_date = None
        for row in ws.iter_rows(min_row=2, max_row=2, values_only=True):
            cell = row[1] if len(row) > 1 else None
            if cell and hasattr(cell, "year"):
                price_date = cell.date() if hasattr(cell, "date") else None
            break

        items: List[SupplierProduct] = []
        for row in ws.iter_rows(min_row=self.DATA_FROM, values_only=True):
            article = clean_str(row[0])
            name = clean_str(row[1])
            if not article or not name:
                continue
            # Пропускаем строки без цены
            price_with_vat = clean_price(row[14]) if len(row) > 14 else None
            price_opt = clean_price(row[16]) if len(row) > 16 else None
            if price_with_vat is None and price_opt is None:
                continue
            # price_base = оптовая без НДС; если нет — из базовой
            if price_opt is not None:
                price_base = price_excl_vat(price_opt)
            elif price_with_vat is not None:
                price_base = price_excl_vat(price_with_vat)
            else:
                continue
            try:
                items.append(SupplierProduct(
                    supplier_code=self.supplier_code,
                    article=article,
                    name=name,
                    unit=clean_unit(row[7]) if len(row) > 7 else "шт",
                    price_base=price_base,
                    price_with_vat=price_with_vat,
                    pack_qty=clean_int(row[11]) if len(row) > 11 else None,
                    ntin=clean_str(row[8]) if len(row) > 8 else None,
                    category_1=clean_str(row[3]) if len(row) > 3 else None,
                    category_2=clean_str(row[4]) if len(row) > 4 else None,
                    category_3=clean_str(row[5]) if len(row) > 5 else None,
                    brand="IEK",
                    status=clean_str(row[9]) if len(row) > 9 else None,
                    price_date=price_date,
                ))
            except Exception:
                continue

        wb.close()
        return self._dedupe(items)

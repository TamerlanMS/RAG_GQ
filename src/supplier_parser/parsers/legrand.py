"""
Парсер прайса Legrand (файл «Tariff KZ...»).
Лист: «Тариф»
Заголовок: строка 4 (1-based)
Данные: строка 5+

Колонки (0-based):
  0  A  ID Reference  (LG-001190)
  1  B  Reference число (1190)
  2  C  Артикул (001190)        = article
  3  D  Наименование             = name
  4  E  Расширенное наименование = name_full
  5  F  Товарная группа
  6  G  Бренд (Legrand / Bticino)
  8  I  Наименование серии       = category_2
  12 M  Статус
  13 N  Минимальная упаковка     = pack_qty
  14 O  Единица измерения        = unit
  15 P  Тариф для дистрибьюторов без НДС = price_base
  22 W  NTIN
"""
from __future__ import annotations

from typing import List

import openpyxl

from src.supplier_parser.models import SupplierProduct
from src.supplier_parser.normalizer import (
    clean_str, clean_price, clean_unit, clean_int, price_incl_vat,
)
from src.supplier_parser.parsers.base import BaseParser

_SHEET_MAIN = "Тариф"
_SHEET_REMOVED = "Удаленные"
_HEADER_ROW = 4
_DATA_FROM = 5


class LegrandParser(BaseParser):
    supplier_code = "legrand"

    def parse(self, file_path: str) -> List[SupplierProduct]:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

        items: List[SupplierProduct] = []
        price_date = None

        for sheet_name, is_removed in [(_SHEET_MAIN, False), (_SHEET_REMOVED, True)]:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=_DATA_FROM, values_only=True):
                if len(row) < 16:
                    continue
                article = clean_str(row[2])
                name = clean_str(row[3])
                if not article or not name:
                    continue
                price_base = clean_price(row[15])
                if price_base is None:
                    continue
                brand = clean_str(row[6]) or "Legrand"
                try:
                    items.append(SupplierProduct(
                        supplier_code=self.supplier_code,
                        article=article,
                        name=name,
                        name_full=clean_str(row[4]),
                        unit=clean_unit(row[14]),
                        price_base=price_base,
                        price_with_vat=price_incl_vat(price_base),
                        pack_qty=clean_int(row[13]),
                        ntin=clean_str(row[22]) if len(row) > 22 else None,
                        category_1=clean_str(row[5]),
                        category_2=clean_str(row[8]),
                        brand=brand,
                        status="удалён" if is_removed else clean_str(row[12]),
                        price_date=price_date,
                    ))
                except Exception:
                    continue

        wb.close()
        return self._dedupe(items)
